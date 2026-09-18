"""Build a tiny workbook with one month block per layout era.

Used by the importer tests so they do not depend on the 400 KB sample, and
so each quirk (fill runs, repeated names, merges past the month end, a
mislabelled month, notes, closures, unlabelled rows) has one known cell.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

WEEKDAYS = ["S", "S", "M", "T", "W", "TR", "F"]
BERTHS = [
    "North Pier West - 410'",
    "North Pier Face - 75'",
    "North Pier East - 240'",
    "Inner Channel - 55'",
    "South Float West - 90'",
    "South Float East - 90'",
]
BLUE = PatternFill(patternType="solid", fgColor="FF00B0F0")
WHITE = PatternFill(patternType="solid", fgColor="FFFFFFFF")


def _era_a_block(ws, top: int, label: str, days: int, extra_rows=()):
    """'AUGUST 1998' in col A; weekday letters; day 1 in B then '=SUM(B3+1)' formulas."""
    ws.cell(top, 1, label)
    for i in range(days):
        ws.cell(top + 1, 2 + i, WEEKDAYS[i % 7])
    ws.cell(top + 2, 2, 1)
    for i in range(1, days):
        prev = openpyxl.utils.get_column_letter(1 + i)
        ws.cell(top + 2, 2 + i, f"=SUM({prev}{top + 2}+1)")
    for j, name in enumerate(list(BERTHS) + list(extra_rows)):
        ws.cell(top + 3 + j, 1, name)
    return top + 3  # first berth row


def _era_b_block(ws, top: int, label: str, days: int):
    """title rows already written; 'JANUARY 2009' with day numbers on the same row."""
    ws.cell(top, 1, label)
    for i in range(days):
        ws.cell(top, 2 + i, i + 1)
        ws.cell(top + 1, 2 + i, WEEKDAYS[i % 7])
    for j, name in enumerate(BERTHS):
        ws.cell(top + 2 + j, 1, name)
    return top + 2


def _era_c_block(ws, top: int, label: str, days: int):
    """'January' with day numbers starting in column C, then the group rows."""
    ws.cell(top, 1, label)
    for i in range(days):
        ws.cell(top, 3 + i, i + 1)
        ws.cell(top + 1, 3 + i, WEEKDAYS[i % 7])
    rows = BERTHS + ["North Finger Piers:", "Small craft slips (institution boats)"]
    for j, name in enumerate(rows):
        ws.cell(top + 2 + j, 1, name)
    return top + 2


def build(path: str | Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- era A: 1998, August and September
    ws = wb.create_sheet("1998")
    r = _era_a_block(ws, 1, "AUGUST 1998", 31)
    ws.cell(r, 3, "F/V Swift Dory")  # Aug 2, fill run C..E -> Aug 2-4
    for c in (3, 4, 5):
        ws.cell(r, c).fill = BLUE
    ws.cell(r, 6).fill = WHITE  # white must not extend the run
    ws.cell(r + 2, 7, "S/V FAR HORIZON")  # Aug 6-7, repeated name
    ws.cell(r + 2, 8, "S/V FAR HORIZON")
    ws.cell(r + 4, 11, "ETA 1200")  # a note, not a booking
    ws.cell(r, 32, "R/V Long Stay")  # Aug 31, continues into September
    r2 = _era_a_block(ws, r + 10, "SEPTEMBER 1998", 30)
    ws.cell(r2, 2, "R/V Long Stay")  # Sep 1
    ws.cell(r2 + 8, 5, "Stray text in an unlabelled row")

    # ---- era B: 2009 with title rows, merges, and a merge past the month end
    ws = wb.create_sheet("2009")
    ws.cell(1, 1, "Harborview Marine Research Center")
    ws.cell(2, 1, "2009 Pier & Dock Schedule (synthetic sample data)")
    ws.cell(3, 1, "Contact: Dock Coordinator")
    r = _era_b_block(ws, 5, "JANUARY 2009", 31)
    ws.cell(r, 3, "R/V GOLDEN COMPASS")  # Jan 2-5 merged
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ws.cell(r + 2, 30, "M/V NORTHERN HARBOR")  # Jan 29.. merged to AH (2 past Jan 31)
    ws.merge_cells(start_row=r + 2, start_column=30, end_row=r + 2, end_column=34)
    ws.cell(r + 1, 10, "Bollard replacement, west face")  # closure on the Face, Jan 9-11 merged
    ws.merge_cells(start_row=r + 1, start_column=10, end_row=r + 1, end_column=12)

    # ---- era B, mislabelled: a 2010 sheet whose second block says 2018
    ws = wb.create_sheet("2010")
    r = _era_b_block(ws, 1, "OCTOBER 2010", 31)
    ws.cell(r, 2, "OSV Clear Osprey")
    # the real 2010 sheet: the true "1 2 3 4 5 6" sits on the row ABOVE the header (day 1 in D),
    # the header's first cells hold vessel names, and its day numbers resume at 7 one column early
    for i in range(6):
        ws.cell(11, 4 + i, i + 1)
    ws.cell(12, 1, "NOVEMBER 2018")
    ws.cell(12, 2, "F/V Junk One")
    ws.cell(12, 3, "M/V Junk Two")
    for i, day in enumerate(range(7, 31)):
        ws.cell(12, 9 + i, day)
    for i in range(30):
        ws.cell(13, 4 + i, WEEKDAYS[i % 7])
    for j, name in enumerate(BERTHS):
        ws.cell(14 + j, 1, name)
    ws.cell(17, 6, "Tug Blue Fathom")  # column F = day 3 by the explicit row (day 4 by the header)

    # ---- era C: 2014 with group rows, an event, a duplicate berth row
    ws = wb.create_sheet("2014")
    ws.cell(1, 1, "Harborview Marine Research Center")
    ws.cell(2, 1, "2014 Pier & Dock Schedule (synthetic sample data)")
    ws.cell(3, 1, "Contact: Dock Coordinator")
    r = _era_c_block(ws, 5, "January", 31)
    ws.cell(r + 5, 6, "Community sail day")  # South Float East, Jan 4-5 merged
    ws.merge_cells(start_row=r + 5, start_column=6, end_row=r + 5, end_column=7)
    ws.cell(r + 6, 3, "F/V Tiny One")  # finger piers Jan 1-3 merged
    ws.merge_cells(start_row=r + 6, start_column=3, end_row=r + 6, end_column=5)
    ws.cell(r + 8, 1, "South Float East - 90'")  # a second row for the same berth
    ws.cell(r + 8, 6, "M/V Second Boat")

    # ---- registry with lengths for two of the grid vessels
    ws = wb.create_sheet("Science")
    ws.append(["VESSEL", "OPERATOR", "CONTACT", "WORK#", "CELL#", "EMAIL", "NOTES"])
    ws.append(["R/V Golden Compass 120'", "Harbor Institute", "Capt. Dana Everly", "Cell: 555-0102"])
    ws.append(["", "Cell: 555-0103"])
    ws.append(["F/V Swift Dory 32'", "", "", "", "", "", "Will raft alongside if needed"])
    ws = wb.create_sheet("Yachts")
    ws.append(["M/Y Long Dory 52'", "LOA: 65', Draft: 4'"])

    ws = wb.create_sheet("8YR Dock Summary")
    ws.append(["8 YR Dock Usage", 2009, 2010])
    ws.append(["North Pier West", 4, 1])
    ws.append(["Total Days", "=SUM(B2:B2)", "=SUM(C2:C2)"])

    ws = wb.create_sheet("Tours")
    ws.append(["Tours are now tracked in a separate workbook; this tab is kept for reference."])
    ws.append(["Date", "Time", "Guide", "Guest", "People", "Dock/ Ship", "Notes"])
    ws.append([43219, 1530, "Avery", "Finley Ingram (Regional Fisheries Agency)", "~6", "R/V Silver Tern", "Confirmed"])
    ws.append([43222, "tbd", "Jesse", "Morgan Ransom (Harbor Institute)", 4, "R/V Blue Heron", ""])
    ws.append(["Requires shore power"])  # a stray fragment with no date

    path = Path(path)
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys

    print(build(sys.argv[1] if len(sys.argv) > 1 else "tiny_schedule.xlsx"))
