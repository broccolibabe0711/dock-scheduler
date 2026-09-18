"""Turn the legacy workbook into berths, vessels, reservations and issues.

The 23 year sheets are grids: a stack of month blocks, one row per berth,
one column per day, the vessel or event name written in the first day's
cell. Three eras use different headers, and how a multi-day stay is shown
changed over time (merged cells from 2009; fills, borders or repeated names
before that). Everything the importer cannot read with confidence becomes
an ImportIssue instead of a guess.

Reading order for one cell: is it the top-left of a merged range (span =
the merge)? else is it coloured, with empty cells of the same colour to the
right (span = the run)? else is the same text repeated in the next cells
(span = the repeat)? else one day. Each reservation remembers which rule
produced its end date in `legacy_ref`.
"""
from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from .classify import CellKind, classify_cell
from .models import Berth, CapacityMode, DayRange, Reservation, ReservationKind, Vessel, name_key
from .registry import read_registry, registry_vessels

MONTHS = [m.lower() for m in calendar.month_name[1:]]
MONTH_RE = re.compile(r"^\s*(" + "|".join(MONTHS) + r")\b\s*(\d{4})?\s*$", re.I)
WEEKDAY_RE = re.compile(r"^(M|T|W|TR|TH|F|S|SU|SA|MO|TU|WE|FR)$")
GROUP_RE = re.compile(r"^(North Finger Piers|Small craft slips)", re.I)
BERTH_LABEL_RE = re.compile(r"^(.*?)\s*-\s*(\d+(?:\.\d+)?)\s*'\s*$")
YEAR_SHEET_RE = re.compile(r"^(19|20)\d\d$")
CANONICAL_BERTHS = (
    "North Pier West",
    "North Pier Face",
    "North Pier East",
    "Inner Channel",
    "South Float West",
    "South Float East",
)
SPAN_MERGE, SPAN_FILL, SPAN_REPEAT, SPAN_SINGLE = "merge", "fill_run", "repeat", "single"


@dataclass
class ImportIssue:
    """Something the importer noticed and did not silently fix."""

    kind: str
    severity: str  # error | warning | info
    sheet: str
    cell: str
    message: str


@dataclass
class Annotation:
    """Text in a berth's cells that is not a booking: 'ETA 1200', 'Fuel truck'."""

    berth_id: int | None
    day: date | None
    text: str
    label: str
    legacy_ref: str


@dataclass
class Tour:
    day: date | None
    time_raw: str
    guide: str
    guest: str
    organisation: str
    people: int | None
    approximate: bool
    vessel_name: str
    notes: str


@dataclass
class RawStay:
    """One run of cells in a berth row, before classification."""

    sheet: str
    block: str
    block_year: int
    block_month: int
    berth_label: str
    row: int
    col: int
    start_day: int
    end_day: int
    last_day: int
    source: str
    text: str


@dataclass
class ImportResult:
    berths: list[Berth] = field(default_factory=list)
    vessels: list[Vessel] = field(default_factory=list)
    reservations: list[Reservation] = field(default_factory=list)
    annotations: list[Annotation] = field(default_factory=list)
    tours: list[Tour] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    raw_stays: list[RawStay] = field(default_factory=list)
    stats: dict[str, Any] = field(default_factory=dict)
    usage_summary: dict[tuple[str, int], int] = field(default_factory=dict)  # from the 8YR sheet

    def issue(self, kind: str, severity: str, sheet: str, cell: str, message: str) -> None:
        self.issues.append(ImportIssue(kind, severity, sheet, cell, message))

    def issue_counts(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for i in self.issues:
            counts[i.kind] = counts.get(i.kind, 0) + 1
        return dict(sorted(counts.items(), key=lambda kv: -kv[1]))


# ---------------------------------------------------------------- reading a sheet
def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_int(v: Any) -> bool:
    return _is_number(v) and float(v).is_integer()


def _text(v: Any) -> str:
    return " ".join(v.split()) if isinstance(v, str) else ""


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _fill_signature(cell) -> str:
    """A short string naming the cell's background, or '' for no colour.

    White and the theme's default background count as no colour: the grid
    paints those on empty cells too, so they cannot mark a stay.
    """
    fill = cell.fill
    pattern = getattr(fill, "patternType", None)
    if not pattern or pattern == "none":
        return ""
    color = fill.fgColor
    if color is None:
        return ""
    if color.type == "rgb":
        rgb = color.rgb if isinstance(color.rgb, str) else ""
        if rgb.upper() in ("FFFFFFFF", "00FFFFFF"):
            return ""
        return f"{pattern}:rgb={rgb}"
    if color.type == "theme":
        if color.theme == 0 and not color.tint:
            return ""
        return f"{pattern}:theme={color.theme},tint={color.tint}"
    if color.type == "indexed":
        return f"{pattern}:indexed={color.indexed}"
    return f"{pattern}:{color.type}"


class _Grid:
    """A worksheet read once into values, fills and merged ranges."""

    def __init__(self, ws) -> None:
        self.name = ws.title
        self.max_row = ws.max_row
        self.max_col = ws.max_column
        self._values: dict[tuple[int, int], Any] = {}
        self._fills: dict[tuple[int, int], str] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    self._values[(cell.row, cell.column)] = cell.value
                sig = _fill_signature(cell)
                if sig:
                    self._fills[(cell.row, cell.column)] = sig
        self.merge_at: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for rng in ws.merged_cells.ranges:
            self.merge_at[(rng.min_row, rng.min_col)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)

    def value(self, r: int, c: int) -> Any:
        return self._values.get((r, c))

    def text(self, r: int, c: int) -> str:
        return _text(self.value(r, c))

    def fill(self, r: int, c: int) -> str:
        return self._fills.get((r, c), "")

    def ref(self, r: int, c: int) -> str:
        return f"{get_column_letter(c)}{r}"


@dataclass
class _DayRow:
    row: int
    day1_col: int
    first: int
    last: int
    length: int


def _day_run(grid: _Grid, r: int) -> _DayRow | None:
    """The longest run of consecutive integers (or '=...+1' formulas) in one row."""
    best: _DayRow | None = None
    c = 2
    while c <= grid.max_col:
        v = grid.value(r, c)
        if _is_int(v):
            start, first = c, int(v)
            expect = last = first
            while c <= grid.max_col:
                v2 = grid.value(r, c)
                if _is_int(v2) and int(v2) == expect:
                    last, expect, c = expect, expect + 1, c + 1
                elif _is_formula(v2) and re.search(r"\+\s*1\)?$", v2):
                    last, expect, c = expect, expect + 1, c + 1
                else:
                    break
            length = last - first + 1
            if best is None or length > best.length:
                best = _DayRow(r, start - (first - 1), first, last, length)
        else:
            c += 1
    return best if best is not None and best.length >= 3 else None


def _find_day_row(grid: _Grid, header_row: int, lower_bound: int, result: ImportResult) -> _DayRow | None:
    """The row holding the day numbers.

    Usually the header row or one of the two below it. In the two damaged
    2010 blocks the header's first day cells were overwritten with vessel
    names and the real "1 2 3 ..." sits on the row ABOVE the header, so that
    row is scanned too (never past the previous block). A run that starts at
    1 is trusted over a longer run that only implies where 1 would be, and
    when candidate rows disagree about the day-1 column the choice is logged.
    """
    candidates: list[_DayRow] = []
    for r in range(max(header_row - 1, lower_bound), min(header_row + 2, grid.max_row) + 1):
        run = _day_run(grid, r)
        if run is not None:
            candidates.append(run)
    if not candidates:
        return None
    explicit = [c for c in candidates if c.first == 1]
    pool = explicit or candidates
    chosen = max(pool, key=lambda c: (c.length, -abs(c.row - header_row)))
    chosen = _DayRow(chosen.row, chosen.day1_col, chosen.first, max(c.last for c in candidates), chosen.length)
    others = {c.day1_col for c in candidates if c.day1_col != chosen.day1_col}
    if others:
        where = "; ".join(f"row {c.row} (days {c.first}..{c.last}, day 1 at {get_column_letter(c.day1_col)})" for c in candidates)
        result.issue("day_row_conflict", "warning", grid.name, grid.ref(chosen.row, chosen.day1_col),
                     f"rows disagree about where day 1 is: {where}; using row {chosen.row}")
    return chosen


def _parse_berth_label(label: str) -> tuple[str, float | None]:
    m = BERTH_LABEL_RE.match(label)
    if m:
        return m.group(1).strip(), float(m.group(2))
    return re.sub(r":\s*$", "", label.strip()), None


def _days_in_month(year: int, month: int) -> int:
    return calendar.monthrange(year, month)[1]


@dataclass
class _Block:
    """One month block of a year sheet, once its header has been understood."""

    sheet: str
    label: str
    year: int
    month: int
    day1_col: int
    header_days: int  # how many days the header lists
    cal_days: int  # how many the calendar has; the calendar wins
    header_rows: list[int]

    @property
    def last_col(self) -> int:
        return self.day1_col + max(self.header_days, self.cal_days) - 1


def _read_year_sheet(grid: _Grid, sheet_year: int, result: ImportResult) -> None:
    """Walk one year sheet block by block and collect raw stays."""
    sheet = grid.name
    stats = {"blocks": [], "candidates": 0, "records": 0, "sources": {}, "outside": 0}
    result.stats[sheet] = stats
    header_rows_all = [r for r in range(1, grid.max_row + 1) if MONTH_RE.match(grid.text(r, 1))]

    for k, h in enumerate(header_rows_all):
        stop = header_rows_all[k + 1] if k + 1 < len(header_rows_all) else grid.max_row + 1
        lower_bound = header_rows_all[k - 1] + 1 if k > 0 else 1
        label = grid.text(h, 1)
        m = MONTH_RE.match(label)
        assert m is not None
        block_month = MONTHS.index(m.group(1).lower()) + 1
        block_year = int(m.group(2)) if m.group(2) else sheet_year
        if block_year > sheet_year:
            # the 2010 sheet ends with "NOVEMBER 2018" / "DECEMBER 2018"; they follow October 2010
            result.issue("mislabelled_month", "warning", sheet, grid.ref(h, 1),
                         f"block '{label}' sits on the {sheet_year} sheet; imported as {sheet_year}-{block_month:02d}")
            block_year = sheet_year
        cal_days = _days_in_month(block_year, block_month)

        day_row = _find_day_row(grid, h, lower_bound, result)
        if day_row is not None:
            day1_col, header_days, day_row_index = day_row.day1_col, day_row.last, day_row.row
            if day_row.first != 1:
                result.issue("day_row_not_from_1", "info", sheet, grid.ref(day_row.row, 1),
                             f"block '{label}' numbers its days {day_row.first}..{day_row.last}; day 1 placed at column {get_column_letter(day1_col)}")
        else:
            day1_col, header_days, day_row_index = (3 if m.group(2) is None else 2), cal_days, None
            result.issue("day_row_missing", "warning", sheet, grid.ref(h, 1),
                         f"block '{label}' has no row of day numbers; day 1 assumed at column {get_column_letter(day1_col)}")
        if header_days != cal_days:
            result.issue("header_days_mismatch", "warning", sheet, grid.ref(h, 1),
                         f"block '{label}' lists {header_days} days but the month has {cal_days}; the calendar wins")
        stats["blocks"].append(f"{label} [day1={get_column_letter(day1_col)}, {header_days}d]")

        header_rows = []
        for r in (h, h + 1, h + 2):
            if r >= stop or r > grid.max_row:
                continue
            is_weekday_row = grid.text(r, 1) == "" and sum(
                1 for c in range(2, grid.max_col + 1) if WEEKDAY_RE.match(grid.text(r, c))
            ) >= 5
            if r == h or r == day_row_index or is_weekday_row:
                header_rows.append(r)
        block = _Block(sheet, label, block_year, block_month, day1_col, header_days, cal_days, header_rows)
        for r in header_rows:
            for c in range(2, grid.max_col + 1):
                v = grid.value(r, c)
                t = _text(v)
                if v is None or _is_number(v) or _is_formula(v) or not t or WEEKDAY_RE.match(t):
                    continue
                result.issue("header_junk", "warning", sheet, grid.ref(r, c),
                             f"'{t}' written in a header row of block '{label}'; not imported")

        group_label = ""
        seen_labels: dict[str, list[int]] = {}
        for r in range(h + 1, stop):
            if r in header_rows:
                continue
            a = grid.text(r, 1)
            if not a:
                for c in range(2, grid.max_col + 1):
                    v = grid.value(r, c)
                    t = _text(v)
                    if v is None or _is_number(v) or _is_formula(v) or not t or WEEKDAY_RE.match(t):
                        continue
                    result.issue("unlabelled_row_text", "info", sheet, grid.ref(r, c),
                                 f"'{t}' in a row with no berth label under block '{label}'; not imported")
                continue
            seen_labels.setdefault(a, []).append(r)
            berth_name, _ = _parse_berth_label(a)
            if GROUP_RE.match(a):
                if a.lower().startswith("north finger piers") or not group_label:
                    group_label = a
            elif berth_name not in CANONICAL_BERTHS:
                result.issue("unexpected_berth_label", "warning", sheet, grid.ref(r, 1),
                             f"row label '{a}' is not one of the known berths")
            _read_berth_row(grid, r, a, block, stats, result)
        for lab, rows in seen_labels.items():
            if len(rows) > 1:
                result.issue("duplicate_berth_row", "info", sheet, grid.ref(rows[1], 1),
                             f"'{lab}' appears on rows {rows} of block '{label}': two occupants written on separate rows")


def _read_berth_row(grid: _Grid, r: int, berth_label: str, block: _Block, stats: dict, result: ImportResult) -> None:
    """Turn the cells of one berth row into raw stays, one per run of cells."""
    sheet, last_col = block.sheet, block.last_col
    c = 2
    while c <= grid.max_col:
        v = grid.value(r, c)
        if v is None or _is_formula(v) or not _text(v):
            c += 1
            continue
        if _is_number(v):
            result.issue("bare_number", "info", sheet, grid.ref(r, c), f"number {v} in a berth row; skipped")
            c += 1
            continue
        text = _text(v)
        stats["candidates"] += 1
        start_day = c - block.day1_col + 1
        fill = grid.fill(r, c)
        end_col, source = c, SPAN_SINGLE
        merge = grid.merge_at.get((r, c))
        if merge is not None:
            source = SPAN_MERGE
            _, _, max_row, max_col = merge
            end_col = max_col
            if max_row != r:
                result.issue("multi_row_merge", "warning", sheet, grid.ref(r, c),
                             f"merged range for '{text}' spans rows {r}..{max_row}; read on row {r} only")
            if end_col > last_col:
                result.issue("merge_beyond_month", "warning", sheet, grid.ref(r, c),
                             f"'{text}' is merged {end_col - last_col} column(s) past the last day of '{block.label}'; cut at the month end")
                end_col = max(last_col, c)
        elif fill:
            e = c
            while e + 1 <= last_col and grid.text(r, e + 1) == "" and grid.fill(r, e + 1) == fill:
                e += 1
            if e > c:
                source, end_col = SPAN_FILL, e
        if source == SPAN_SINGLE:
            e = c
            while e + 1 <= last_col and grid.value(r, e + 1) == v:
                e += 1
            if e > c:
                source, end_col = SPAN_REPEAT, e
        next_c = end_col + 1
        if start_day < 1 or start_day > block.cal_days:
            stats["outside"] += 1
            where = "before day 1" if start_day < 1 else f"past the {block.cal_days} days of the month"
            result.issue("cell_outside_day_columns", "warning", sheet, grid.ref(r, c),
                         f"'{text}' sits {where} of '{block.label}' (day {start_day}); not imported")
            c = next_c
            continue
        if start_day > block.header_days:
            result.issue("beyond_header_days", "info", sheet, grid.ref(r, c),
                         f"'{text}' is on day {start_day}, which the header of '{block.label}' does not list but the calendar has; imported")
        end_day = min(end_col - block.day1_col + 1, block.cal_days)
        result.raw_stays.append(RawStay(sheet, block.label, block.year, block.month, berth_label, r, c,
                                        start_day, end_day, block.header_days, source, text))
        stats["records"] += 1
        stats["sources"][source] = stats["sources"].get(source, 0) + 1
        c = next_c


# ---------------------------------------------------------------- assembling the model
def _berth_for_label(label: str, first_year: int) -> Berth:
    name, length = _parse_berth_label(label)
    if GROUP_RE.match(label):
        # a group of slips drawn as one row: how many boats it holds is unknown,
        # so it is a linear berth with no length and shared days come out "unverifiable"
        return Berth(name, None, CapacityMode.LINEAR, active_from=date(first_year, 1, 1))
    return Berth(name, length, CapacityMode.LINEAR)


def _stitch(reservations: list[Reservation], result: ImportResult) -> list[Reservation]:
    """Join stays the grid split at a month boundary into one reservation.

    Same berth, same kind, same vessel or title, the first ending on the last
    day of its month and the next starting the following day."""
    def identity(r: Reservation):
        return (r.berth_id, r.kind, r.vessel.key if r.vessel else r.title.lower())

    by_identity: dict[tuple, list[Reservation]] = {}
    for r in reservations:
        by_identity.setdefault(identity(r), []).append(r)
    out: list[Reservation] = []
    stitched = 0
    for group in by_identity.values():
        group.sort(key=lambda r: (r.days.start, r.days.end))
        current = group[0]
        for nxt in group[1:]:
            month_end = current.days.end.day == _days_in_month(current.days.end.year, current.days.end.month)
            if month_end and nxt.days.start == current.days.end + timedelta(days=1):
                current = Reservation(
                    berth_id=current.berth_id, kind=current.kind,
                    days=DayRange(current.days.start, max(current.days.end, nxt.days.end)),
                    vessel=current.vessel, title=current.title, status=current.status,
                    source=current.source, legacy_ref=f"{current.legacy_ref}+{nxt.legacy_ref}",
                    notes="stitched across a month boundary",
                )
                stitched += 1
            else:
                out.append(current)
                current = nxt
        out.append(current)
    result.stats["stitched"] = stitched
    return out


def _read_usage_summary(wb, result: ImportResult) -> None:
    if "8YR Dock Summary" not in wb.sheetnames:
        return
    rows = list(wb["8YR Dock Summary"].iter_rows(values_only=True))
    years = [int(y) for y in rows[0][1:] if _is_int(y)]
    for row in rows[1:]:
        name = _text(row[0])
        if not name or name.lower().startswith("total"):
            continue
        for y, v in zip(years, row[1:]):
            if _is_int(v):
                result.usage_summary[(name, y)] = int(v)


def _read_tours(wb, result: ImportResult) -> None:
    if "Tours" not in wb.sheetnames:
        return
    for row_index, row in enumerate(wb["Tours"].iter_rows(values_only=True), start=1):
        cells = list(row) + [None] * 7
        raw_date, raw_time, guide, guest, people, vessel, notes = cells[:7]
        texts = [_text(c) for c in cells[:7] if c is not None and _text(c)]
        if isinstance(raw_date, datetime):
            day = raw_date.date()
        elif _is_number(raw_date):
            day = date(1899, 12, 30) + timedelta(days=int(raw_date))  # Excel's 1900 date system
        else:
            if texts and _text(raw_date).lower() != "date" and "separate workbook" not in " ".join(texts).lower():
                result.issue("unreadable_tour_row", "info", "Tours", f"A{row_index}",
                             f"row has no readable date: {' | '.join(texts)}; not imported")
            continue  # banner or header rows carry no tour
        guest_text = _text(guest)
        m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", guest_text)
        people_text = _text(people) if not _is_number(people) else str(int(people))
        digits = re.sub(r"[^0-9]", "", people_text)
        result.tours.append(Tour(
            day=day, time_raw=_text(raw_time) if not _is_number(raw_time) else str(int(raw_time)),
            guide=_text(guide), guest=m.group(1) if m else guest_text,
            organisation=m.group(2) if m else "", people=int(digits) if digits else None,
            approximate=people_text.startswith("~"), vessel_name=_text(vessel), notes=_text(notes),
        ))


def import_workbook(path: str | Path) -> ImportResult:
    """Read the whole legacy workbook. Never raises on bad data; look at .issues."""
    wb = openpyxl.load_workbook(path)
    result = ImportResult()

    # 1. registries: the only source of vessel lengths
    registry_issues: list[str] = []
    entries = read_registry(wb, registry_issues)
    registry, conflicts = registry_vessels(entries)
    for text in registry_issues:
        result.issue("registry_row_unplaced", "warning", "Science/Yachts", "", text)
    for text in conflicts:
        result.issue("registry_length_conflict", "warning", "Science/Yachts", "", text)
    vessel_by_key: dict[str, Vessel] = {v.key: v for v in registry}

    # 2. the grids
    year_sheets = sorted(n for n in wb.sheetnames if YEAR_SHEET_RE.match(n))
    for name in year_sheets:
        _read_year_sheet(_Grid(wb[name]), int(name), result)

    # 3. skip month blocks that appear twice (a sheet that starts with the previous December)
    seen_blocks: dict[tuple[int, int], str] = {}
    kept: list[RawStay] = []
    skipped_blocks: set[tuple[str, str]] = set()
    for stay in result.raw_stays:
        key = (stay.block_year, stay.block_month)
        owner = seen_blocks.setdefault(key, stay.sheet)
        if owner != stay.sheet:
            skipped_blocks.add((stay.sheet, stay.block))
            continue
        kept.append(stay)
    for sheet, block in sorted(skipped_blocks):
        n = sum(1 for s in result.raw_stays if s.sheet == sheet and s.block == block)
        result.issue("duplicate_block", "warning", sheet, "", f"block '{block}' repeats a month already read from another sheet; its {n} cell(s) were not imported")

    # 4. berths, in the order the grids introduce them
    first_year: dict[str, int] = {}
    for stay in kept:
        first_year.setdefault(stay.berth_label, stay.block_year)
        first_year[stay.berth_label] = min(first_year[stay.berth_label], stay.block_year)
    def berth_order(item: tuple[str, int]) -> tuple:
        label, year = item
        name = _parse_berth_label(label)[0]
        if name in CANONICAL_BERTHS:
            return (0, CANONICAL_BERTHS.index(name), label)
        return (1, year, label)  # group rows after the six pier faces, oldest first

    berths: dict[str, Berth] = {}
    for i, (label, year) in enumerate(sorted(first_year.items(), key=berth_order), start=1):
        b = _berth_for_label(label, year)
        berths[label] = Berth(b.name, b.length_ft, b.capacity_mode, b.clearance_ft, b.active_from, b.active_to, id=i)
    # the same berth name may carry two labels over the years; map by name
    berth_by_name: dict[str, Berth] = {}
    for label, b in berths.items():
        berth_by_name.setdefault(b.name, b)
    result.berths = list(berth_by_name.values())

    # 5. classify every stay
    vessels: dict[str, Vessel] = {}
    reservations: list[Reservation] = []
    for stay in kept:
        ref = f"{stay.sheet}!{get_column_letter(stay.col)}{stay.row}:{stay.source}"
        cls = classify_cell(stay.text)
        berth = berth_by_name[_parse_berth_label(stay.berth_label)[0]]
        try:
            start = date(stay.block_year, stay.block_month, stay.start_day)
            end = date(stay.block_year, stay.block_month, max(stay.start_day, stay.end_day))
        except ValueError:
            result.issue("invalid_date", "warning", stay.sheet, f"{get_column_letter(stay.col)}{stay.row}",
                         f"'{stay.text}' in '{stay.block}' has no valid date (day {stay.start_day}); not imported")
            continue
        if cls.kind in (CellKind.NOTE, CellKind.OTHER):
            result.annotations.append(Annotation(berth.id, start, stay.text, cls.label, ref))
            if cls.kind is CellKind.OTHER:
                result.issue("unclassified_text", "warning", stay.sheet, f"{get_column_letter(stay.col)}{stay.row}",
                             f"'{stay.text}' did not match any known vessel, event, closure or note; kept as an annotation")
            continue
        if cls.kind is CellKind.VESSEL:
            key = name_key(cls.label)
            vessel = vessel_by_key.get(key) or vessels.get(key) or Vessel(cls.label, None, cls.type_prefix)
            vessels[key] = vessel
            reservations.append(Reservation(berth.id, ReservationKind.VESSEL, DayRange(start, end), vessel=vessel,
                                            status="confirmed", source="import", legacy_ref=ref))
        else:
            kind = ReservationKind.EVENT if cls.kind is CellKind.EVENT else ReservationKind.CLOSURE
            reservations.append(Reservation(berth.id, kind, DayRange(start, end), title=cls.label,
                                            status="confirmed", source="import", legacy_ref=ref))

    # 6. vessels get ids; registry vessels that never appear in the grid are kept too
    all_vessels = list(vessels.values()) + [v for v in registry if v.key not in vessels]
    with_ids: dict[str, Vessel] = {}
    for i, v in enumerate(sorted(all_vessels, key=lambda v: v.name.lower()), start=1):
        with_ids[v.key] = Vessel(v.name, v.length_ft, v.type_prefix, v.draft_ft, v.operator, v.rafts_ok, v.notes, id=i)
    result.vessels = list(with_ids.values())
    reservations = [
        Reservation(r.berth_id, r.kind, r.days, vessel=with_ids[r.vessel.key] if r.vessel else None, title=r.title,
                    status=r.status, source=r.source, legacy_ref=r.legacy_ref, notes=r.notes)
        for r in reservations
    ]

    # 7. stitch month-split stays, then number everything
    reservations = _stitch(reservations, result)
    reservations.sort(key=lambda r: (r.days.start, r.berth_id, r.display_name))
    result.reservations = [
        Reservation(r.berth_id, r.kind, r.days, vessel=r.vessel, title=r.title, status=r.status,
                    source=r.source, legacy_ref=r.legacy_ref, notes=r.notes, id=i)
        for i, r in enumerate(reservations, start=1)
    ]

    _read_usage_summary(wb, result)
    _read_tours(wb, result)
    result.stats["totals"] = {
        "raw_stays": len(result.raw_stays),
        "kept_stays": len(kept),
        "reservations": len(result.reservations),
        "annotations": len(result.annotations),
        "vessels": len(result.vessels),
        "vessels_with_length": sum(1 for v in result.vessels if v.length_ft is not None),
        "berths": len(result.berths),
        "tours": len(result.tours),
        "issues": len(result.issues),
    }
    return result
